import base64
import csv
import io
import json
import logging
from datetime import datetime
from odoo import models, fields, api, _
from markupsafe import Markup
from odoo.exceptions import UserError
import re

_logger = logging.getLogger(__name__)

try:
    import openpyxl
except ImportError:
    _logger.error("OpenPyXL not installed")


def _safe_float(val):
    try:
        if isinstance(val, (int, float)):
            return float(val)
        if isinstance(val, str) and val.strip():
            return float(val.strip().replace('.', '').replace(',', '.')) if ',' in val else float(val.strip())
        return 0.0
    except (TypeError, ValueError):
        return 0.0


def _safe_int(val):
    try:
        if isinstance(val, (int, float)):
            return int(val)
        if isinstance(val, str) and val.strip():
            v = val.strip()
            if v.endswith('.0'):
                v = v[:-2]
            v = v.replace('.', '').replace(',', '')
            return int(re.sub(r'\D', '', v) or 0)
        return 0
    except (TypeError, ValueError):
        return 0


# Columnas que identifican de forma inequívoca el diseño de registro
# "Libro de Compras" (CSV descargado desde ARCA > Libro de IVA Compras),
# distinto del Excel/CSV de "Mis Comprobantes".
LIBRO_COMPRAS_SIGNATURE_HEADER = 'Denominación Vendedor'

LIBRO_COMPRAS_REQUIRED_HEADERS = [
    'Fecha de Emisión', 'Tipo de Comprobante', 'Punto de Venta', 'Número de Comprobante',
    'Tipo Doc. Vendedor', 'Nro. Doc. Vendedor', 'Denominación Vendedor', 'Importe Total',
    'Moneda Original', 'Tipo de Cambio', 'Importe No Gravado', 'Importe Exento',
    'Crédito Fiscal Computable', 'Importe de Per. o Pagos a Cta. de Otros Imp. Nac.',
    'Importe de Percepciones de Ingresos Brutos', 'Importe de Impuestos Municipales',
    'Importe de Percepciones o Pagos a Cuenta de IVA', 'Importe de Impuestos Internos',
    'Importe Otros Tributos', 'Neto Gravado IVA 0%', 'Neto Gravado IVA 2,5%', 'Importe IVA 2,5%',
    'Neto Gravado IVA 5%', 'Importe IVA 5%', 'Neto Gravado IVA 10,5%', 'Importe IVA 10,5%',
    'Neto Gravado IVA 21%', 'Importe IVA 21%', 'Neto Gravado IVA 27%', 'Importe IVA 27%',
    'Total Neto Gravado', 'Total IVA',
]

# rate -> (columna Neto Gravado, columna Importe IVA)
LIBRO_COMPRAS_TAX_RATE_COLUMNS = [
    (0.0, 'Neto Gravado IVA 0%', None),
    (2.5, 'Neto Gravado IVA 2,5%', 'Importe IVA 2,5%'),
    (5.0, 'Neto Gravado IVA 5%', 'Importe IVA 5%'),
    (10.5, 'Neto Gravado IVA 10,5%', 'Importe IVA 10,5%'),
    (21.0, 'Neto Gravado IVA 21%', 'Importe IVA 21%'),
    (27.0, 'Neto Gravado IVA 27%', 'Importe IVA 27%'),
]

# Columnas de percepciones/otros tributos aditivas (se suman al Importe Total)
# junto a lo Gravado/No Gravado/Exento. 'Crédito Fiscal Computable' queda afuera
# a propósito: es informativo, no es un importe adicional del comprobante.
LIBRO_COMPRAS_OTHER_TRIBUTES_COLUMNS = [
    ('perc_otros_imp_nacionales', 'Importe de Per. o Pagos a Cta. de Otros Imp. Nac.', 'Percepción/Pago a Cta. de Otros Imp. Nacionales'),
    ('perc_iibb', 'Importe de Percepciones de Ingresos Brutos', 'Percepción de Ingresos Brutos'),
    ('imp_municipales', 'Importe de Impuestos Municipales', 'Impuestos Municipales'),
    ('perc_iva', 'Importe de Percepciones o Pagos a Cuenta de IVA', 'Percepción/Pago a Cta. de IVA'),
    ('imp_internos', 'Importe de Impuestos Internos', 'Impuestos Internos'),
    ('otros_tributos', 'Importe Otros Tributos', 'Otros Tributos'),
]

# Nombres exactos (o candidatos) de los impuestos de percepción a buscar en
# account.tax para el 'Libro de Compras'. 'Percepción IVA' tiene dos variantes
# según la alícuota de IVA asociada (ver _analyze_libro_compras).
TAX_NAME_PERC_GANANCIAS = ['Perc Gananc', 'Percepción Ganancias']
TAX_NAME_PERC_IIBB = ['P. IIBB MZA', 'Percepción IIBB Mendoza']


class L10nArImportArcaWizard(models.TransientModel):
    _name = 'l10n_ar.arca.import.wizard'
    _description = 'Asistente de Importación ARCA'

    file_data = fields.Binary(string='Archivo Excel', required=True)
    filename = fields.Char(string='Nombre Archivo')
    import_type = fields.Selection([
        ('in_invoice', 'Facturas Recibidas (Proveedores)'),
        ('out_invoice', 'Facturas Emitidas (Clientes)')
    ], string='Tipo de Importación', required=True, default='in_invoice')

    journal_purchase_id = fields.Many2one(
        'account.journal',
        string='Diario de Compras',
        domain="[('type', '=', 'purchase'), ('l10n_latam_use_documents', '=', True)]"
    )
    journal_sale_id = fields.Many2one(
        'account.journal',
        string='Diario de Ventas',
        domain="[('type', '=', 'sale'), ('l10n_latam_use_documents', '=', True)]"
    )

    line_ids = fields.One2many(
        'l10n_ar.arca.import.line', 'wizard_id', string='Líneas')

    state = fields.Selection([
        ('upload', 'Carga'),
        ('review', 'Revisión'),
        ('done', 'Hecho')
    ], default='upload', string='Estado')

    company_id = fields.Many2one(
        'res.company', string='Compañía', required=True, default=lambda self: self.env.company)

    pre_analysis_message = fields.Char(string='Pre-análisis', readonly=True)
    is_auto_detected = fields.Boolean(default=False)

    can_import = fields.Boolean(compute='_compute_can_import')

    @api.depends('line_ids.status', 'line_ids.to_import')
    def _compute_can_import(self):
        for wizard in self:
            wizard.can_import = any(
                l.status == 'ready' and l.to_import for l in wizard.line_ids)

    detected_period = fields.Char(
        string='Período Detectado', compute='_compute_detected_period')

    @api.depends('line_ids', 'line_ids.date')
    def _compute_detected_period(self):
        for wizard in self:
            dates = wizard.line_ids.mapped('date')
            if not dates:
                wizard.detected_period = False
                wizard.period_start = False
                wizard.period_end = False
                wizard.period_month_year = False
                wizard.journal_name_display = False
                wizard.journal_pos_display = False
                continue

            min_date = min(dates)
            max_date = max(dates)

            # Format dates DD/MM/AAAA
            d_from = min_date.strftime('%d/%m/%Y')
            d_to = max_date.strftime('%d/%m/%Y')

            # Get unique months (Spanish names if possible, but standard Babel/Python loc is safer or just specific mapping)
            # Simple approach: "Month Year"
            months = sorted(list(set(d.strftime('%m/%Y') for d in dates)))

            # Context-aware month names could be tricky without babel, let's use a simple mapping or Odoo's format_date if available
            # We'll stick to a simple mapping for "Mes/es detectado/s"
            month_names = {
                '01': 'Enero', '02': 'Febrero', '03': 'Marzo', '04': 'Abril', '05': 'Mayo', '06': 'Junio',
                '07': 'Julio', '08': 'Agosto', '09': 'Septiembre', '10': 'Octubre', '11': 'Noviembre', '12': 'Diciembre'
            }

            month_labels = []
            for m_y in months:
                m, y = m_y.split('/')
                month_labels.append(f"{month_names.get(m, m)} {y}")

            months_str = ", ".join(month_labels)

            # Determine active journal
            journal = wizard.journal_purchase_id if wizard.import_type == 'in_invoice' else wizard.journal_sale_id
            journal_name = journal.name if journal else 'No definido'

            pos_info = ""
            if wizard.import_type == 'out_invoice' and journal and journal.l10n_ar_afip_pos_number:
                pos_info = f" (PdV: {journal.l10n_ar_afip_pos_number})"

            wizard.detected_period = f"Diario: {journal_name}{pos_info} | Período detectado según facturas: {d_from} al {d_to} ({months_str})"

            # Populate new granular fields for UI redesign
            wizard.period_start = min_date.strftime('%d/%m')
            wizard.period_end = max_date.strftime('%d/%m/%Y')
            wizard.period_month_year = months_str
            wizard.journal_name_display = journal_name
            wizard.journal_pos_display = str(journal.l10n_ar_afip_pos_number) if journal.l10n_ar_afip_pos_number else False

    period_start = fields.Char(compute='_compute_detected_period')
    period_end = fields.Char(compute='_compute_detected_period')
    period_month_year = fields.Char(compute='_compute_detected_period')
    journal_name_display = fields.Char(compute='_compute_detected_period')
    journal_pos_display = fields.Char(compute='_compute_detected_period')

    # --- Dashboard Fields ---
    view_filter = fields.Selection([
        ('all', 'Todos'),
        ('ready', 'Listos'),
        ('error', 'Errores'),
        ('exists', 'Existentes')
    ], default='all', string="Filtro de Vista")

    analytic_distribution = fields.Json(string='Distribución Analítica')
    analytic_precision = fields.Integer(store=False, default=2)
    focus_holder = fields.Char(string='Focus Trap')
    
    account_id = fields.Many2one('account.account', string='Cuenta Contable', domain="[('deprecated', '=', False)]")
    tax_other_tributes_id = fields.Many2one(
        'account.tax', string='Impuesto Otros Tributos', check_company=True,
        help="Impuesto aplicado a las filas con monto en la columna 'Otros Tributos'"
    )

    cae_status_mode = fields.Selection([
        ('available', 'CAE Disponible / Requerido'),
        ('not_available', 'CAE No Disponible'),
        ('missing', 'Módulo Oculto')
    ], compute='_compute_cae_status_mode')
    
    is_electronic_journal = fields.Boolean(compute='_compute_is_electronic_journal')

    @api.depends('journal_sale_id', 'journal_purchase_id', 'import_type')
    def _compute_is_electronic_journal(self):
        for rec in self:
            journal = rec.journal_sale_id if rec.import_type == 'out_invoice' else rec.journal_purchase_id
            is_electronic = False
            if journal:
                pos_system = getattr(journal, 'l10n_ar_afip_pos_system', False)
                if pos_system == 'WSFE':
                    is_electronic = True
            rec.is_electronic_journal = is_electronic

    @api.depends('company_id', 'import_type')
    def _compute_cae_status_mode(self):
        for rec in self:
            move_model = self.env['account.move']
            has_enterprise = 'l10n_ar_afip_auth_mode' in move_model._fields
            has_community = 'afip_auth_mode' in move_model._fields
            
            if not has_enterprise and not has_community:
                rec.cae_status_mode = 'missing'
                continue

            if rec.import_type == 'out_invoice':
                # Sales always require CAE if AFIP module is installed
                rec.cae_status_mode = 'available'
            else:
                # Purchases use Verification Setting if Enterprise, else assumed available for Community
                if has_enterprise and hasattr(self.env.company, 'l10n_ar_afip_verification_type'):
                    verif_type = self.env.company.l10n_ar_afip_verification_type
                    rec.cae_status_mode = 'not_available' if verif_type == 'not_available' else 'available'
                else:
                    rec.cae_status_mode = 'available'
    
    # --- Dynamic Summary Fields ---
    company_currency_id = fields.Many2one('res.currency', string='Company Currency', 
                                          related='company_id.currency_id', readonly=True)
    summary_facturas = fields.Monetary(string='Facturas', compute='_compute_summaries', currency_field='company_currency_id')
    summary_ncs = fields.Monetary(string='Notas de Crédito', compute='_compute_summaries', currency_field='company_currency_id')
    summary_iva_21 = fields.Monetary(string='IVA 21%', compute='_compute_summaries', currency_field='company_currency_id')
    summary_iva_105 = fields.Monetary(string='IVA 10.5%', compute='_compute_summaries', currency_field='company_currency_id')
    summary_otros_tributos = fields.Monetary(string='Otros Tributos', compute='_compute_summaries', currency_field='company_currency_id')
    
    # Los 6 campos de percepciones/tributos discriminados (los 5 nuevos de
    # 'Libro de Compras' + el histórico 'otros_tributos') se consolidan en un
    # único total a efectos de resumen/validación. El desglose por tipo queda
    # disponible en la línea para una futura diferenciación del tratamiento
    # contable de cada percepción.
    OTROS_TRIBUTOS_FIELD_NAMES = [f[0] for f in LIBRO_COMPRAS_OTHER_TRIBUTES_COLUMNS]

    def _line_tributes_total(self, line):
        return sum(getattr(line, fname) for fname in self.OTROS_TRIBUTOS_FIELD_NAMES)

    @api.depends('visible_line_ids.to_import', 'visible_line_ids.amount_total', 'visible_line_ids.is_refund',
                 'visible_line_ids.iva_21_amount', 'visible_line_ids.iva_105_amount', 'visible_line_ids.otros_tributos',
                 'visible_line_ids.perc_otros_imp_nacionales', 'visible_line_ids.perc_iibb',
                 'visible_line_ids.imp_municipales', 'visible_line_ids.perc_iva', 'visible_line_ids.imp_internos')
    def _compute_summaries(self):
        for rec in self:
            facturas = 0.0
            ncs = 0.0
            iva_21 = 0.0
            iva_105 = 0.0
            otros_tributos = 0.0

            for line in rec.visible_line_ids.filtered('to_import'):
                if line.is_refund:
                    ncs += line.amount_total
                else:
                    facturas += line.amount_total

                iva_21 += line.iva_21_amount
                iva_105 += line.iva_105_amount
                otros_tributos += rec._line_tributes_total(line)

            rec.summary_facturas = facturas
            rec.summary_ncs = ncs
            rec.summary_iva_21 = iva_21
            rec.summary_iva_105 = iva_105
            rec.summary_otros_tributos = otros_tributos

    tax_domain_type = fields.Char(compute='_compute_tax_domain_type')
    has_otros_tributos = fields.Boolean(compute='_compute_has_otros_tributos')

    @api.depends('visible_line_ids.otros_tributos', 'visible_line_ids.perc_otros_imp_nacionales',
                 'visible_line_ids.perc_iibb', 'visible_line_ids.imp_municipales',
                 'visible_line_ids.perc_iva', 'visible_line_ids.imp_internos', 'visible_line_ids.to_import')
    def _compute_has_otros_tributos(self):
        for rec in self:
            rec.has_otros_tributos = any(
                l.to_import and rec._line_tributes_total(l) > 0 for l in rec.visible_line_ids)
            
    @api.depends('import_type')
    def _compute_tax_domain_type(self):
        for rec in self:
            rec.tax_domain_type = 'purchase' if rec.import_type == 'in_invoice' else 'sale'

    info_all = fields.Char(compute='_compute_dashboard_data')
    info_ready = fields.Char(compute='_compute_dashboard_data')
    info_error = fields.Char(compute='_compute_dashboard_data')
    info_exists = fields.Char(compute='_compute_dashboard_data')

    visible_line_ids = fields.Many2many(
        'l10n_ar.arca.import.line', compute='_compute_visible_lines', store=True)

    partner_filter_ids = fields.One2many(
        'l10n_ar.arca.import.partner.filter', 'wizard_id', string='Filtro de Contactos')
    
    missing_partners_count = fields.Integer(string='Contactos Faltantes', compute='_compute_missing_partners')
    @api.depends('partner_filter_ids.partner_id')
    def _compute_missing_partners(self):
        for wizard in self:
            wizard.missing_partners_count = len(wizard.partner_filter_ids.filtered(lambda p: not p.partner_id))

    def _get_action_reopen(self):
        return {
            'type': 'ir.actions.act_window',
            'name': 'Importación de Facturas de ARCA',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
            'context': self.env.context,
        }

    def action_select_all_visible_lines(self):
        for wizard in self:
            lines_to_update = wizard.visible_line_ids.filtered(lambda l: l.status == 'ready')
            lines_to_update.write({'to_import': True})
        return self._get_action_reopen()

    def action_deselect_all_visible_lines(self):
        for wizard in self:
            wizard.visible_line_ids.write({'to_import': False})
        return self._get_action_reopen()

    def action_select_all_partners(self):
        for wizard in self:
            wizard.partner_filter_ids.write({'selected': True})
        return self._get_action_reopen()

    def action_deselect_all_partners(self):
        for wizard in self:
            wizard.partner_filter_ids.write({'selected': False})
        return self._get_action_reopen()

    show_select_all_partners = fields.Boolean(compute='_compute_partner_selection')
    show_deselect_all_partners = fields.Boolean(compute='_compute_partner_selection')

    @api.depends('partner_filter_ids.selected')
    def _compute_partner_selection(self):
        for wizard in self:
            partners = wizard.partner_filter_ids
            if not partners:
                wizard.show_select_all_partners = False
                wizard.show_deselect_all_partners = False
                continue
            
            selected_count = len(partners.filtered('selected'))
            total_count = len(partners)
            
            wizard.show_select_all_partners = selected_count < total_count
            wizard.show_deselect_all_partners = selected_count > 0

    show_select_all_ready = fields.Boolean(compute='_compute_ready_line_selection')
    show_deselect_all_ready = fields.Boolean(compute='_compute_ready_line_selection')

    @api.depends('visible_line_ids', 'visible_line_ids.to_import', 'visible_line_ids.status')
    def _compute_ready_line_selection(self):
        for wizard in self:
            # We need to ensure we are looking at the current visible set
            # filtered by status='ready' as that's what the buttons control
            ready_lines = wizard.visible_line_ids.filtered(lambda l: l.status == 'ready')
            
            if not ready_lines:
                wizard.show_select_all_ready = False
                wizard.show_deselect_all_ready = False
                continue
            
            # Count how many are selected
            selected_count = len(ready_lines.filtered('to_import'))
            total_count = len(ready_lines)
            
            # Show "Select All" if not all are selected (e.g. 0 or some)
            wizard.show_select_all_ready = selected_count < total_count
            
            # Show "Deselect All" if at least one is selected
            wizard.show_deselect_all_ready = selected_count > 0

    @api.depends('line_ids', 'view_filter', 'partner_filter_ids.selected')
    def _compute_visible_lines(self):
        for wizard in self:
            domain = []

            # 1. Status Filter
            if wizard.view_filter == 'ready':
                domain.append(('status', '=', 'ready'))
            elif wizard.view_filter == 'error':
                domain.append(('status', '=', 'error'))
            elif wizard.view_filter == 'exists':
                domain.append(('status', '=', 'exists'))

            lines = wizard.line_ids.filtered_domain(domain)

            # 2. Partner Filter
            selected_partners = wizard.partner_filter_ids.filtered('selected')
            if selected_partners:
                # Filter lines where partner_name matches any of the selected partner names
                selected_names = selected_partners.mapped('name')
                lines = lines.filtered(
                    lambda l: l.partner_name in selected_names)

            wizard.visible_line_ids = lines

    @api.onchange('partner_filter_ids', 'view_filter')
    def _onchange_filter_update(self):
        # Explicitly trigger recompute of visible lines when filters change
        # This helps in wizard context where One2many changes might not auto-propagate to other fields
        self._compute_visible_lines()

    total_lines = fields.Integer(
        compute='_compute_total_lines', string='Total Comprobantes en Archivo')

    @api.depends('line_ids')
    def _compute_total_lines(self):
        for wizard in self:
            wizard.total_lines = len(wizard.line_ids)

    has_selected_partners = fields.Boolean(
        compute='_compute_has_selected_partners', string='Tiene Contactos Seleccionados')

    @api.depends('partner_filter_ids.selected')
    def _compute_has_selected_partners(self):
        for wizard in self:
            wizard.has_selected_partners = bool(wizard.partner_filter_ids.filtered('selected'))

    @api.depends('line_ids', 'line_ids.status', 'partner_filter_ids.selected')
    def _compute_dashboard_data(self):
        for wizard in self:
            selected_partners = wizard.partner_filter_ids.filtered('selected')
            
            # IF NO PARTNERS SELECTED -> COUNTS ARE 0
            if not selected_partners:
                c_all = 0
                c_ready = 0
                c_error = 0
                c_exists = 0
            else:
                # Filtering logic with selection
                selected_names = selected_partners.mapped('name')
                base_lines = wizard.line_ids.filtered(
                    lambda l: l.partner_name in selected_names)

                c_all = len(base_lines)
                c_ready = len(base_lines.filtered(lambda l: l.status == 'ready'))
                c_error = len(base_lines.filtered(lambda l: l.status == 'error'))
                c_exists = len(base_lines.filtered(lambda l: l.status == 'exists'))

            # Format: "LABEL: COUNT"
            # Simplified labels as requested
            wizard.info_all = f"TODOS: {c_all}"
            wizard.info_ready = f"LISTOS: {c_ready}"
            wizard.info_error = f"ERRORES: {c_error}"
            wizard.info_exists = f"EXISTENTES: {c_exists}"

            # New Integer fields for custom UI
            wizard.count_all = c_all
            wizard.count_ready = c_ready
            wizard.count_error = c_error
            wizard.count_exists = c_exists

    count_all = fields.Integer(compute='_compute_dashboard_data')
    count_ready = fields.Integer(compute='_compute_dashboard_data')
    count_error = fields.Integer(compute='_compute_dashboard_data')
    count_exists = fields.Integer(compute='_compute_dashboard_data')



    has_missing_partners_selection = fields.Boolean(
        compute='_compute_has_missing_partners_selection', string='Tiene Contactos Faltantes')

    @api.depends('partner_filter_ids.partner_id', 'partner_filter_ids.selected')
    def _compute_has_missing_partners_selection(self):
        for wizard in self:
            selected = wizard.partner_filter_ids.filtered('selected')
            if not selected:
                wizard.has_missing_partners_selection = False
            else:
                wizard.has_missing_partners_selection = any(not p.partner_id for p in selected)

    @api.depends('line_ids', 'view_filter', 'partner_filter_ids.selected')
    def _compute_visible_lines(self):
        for wizard in self:
            # 1. Partner Filter
            selected_partners = wizard.partner_filter_ids.filtered('selected')
            
            if not selected_partners:
                wizard.visible_line_ids = False
                continue

            selected_names = selected_partners.mapped('name')
            base_lines = wizard.line_ids.filtered(
                lambda l: l.partner_name in selected_names)

            # 2. Status Filter
            if wizard.view_filter == 'all':
                wizard.visible_line_ids = base_lines
            elif wizard.view_filter == 'ready':
                wizard.visible_line_ids = base_lines.filtered(
                    lambda l: l.status == 'ready')
            elif wizard.view_filter == 'error':
                wizard.visible_line_ids = base_lines.filtered(
                    lambda l: l.status == 'error')
            elif wizard.view_filter == 'exists':
                wizard.visible_line_ids = base_lines.filtered(
                    lambda l: l.status == 'exists')
            else:
                wizard.visible_line_ids = base_lines

    def action_set_filter_all(self):
        self.view_filter = 'all'
        return self._reload_view()

    def action_set_filter_ready(self):
        self.view_filter = 'ready'
        return self._reload_view()

    def action_set_filter_error(self):
        self.view_filter = 'error'
        return self._reload_view()

    def action_set_filter_exists(self):
        self.view_filter = 'exists'
        return self._reload_view()



    def _reload_view(self):
        return {
            'type': 'ir.actions.act_window',
            'name': 'Importación de Facturas de ARCA',
            'res_model': self._name,
            'view_mode': 'form',
            'res_id': self.id,
            'target': 'new',
        }

    def _decode_csv_text(self, file_content):
        for encoding in ('utf-8-sig', 'cp1252'):
            try:
                return file_content.decode(encoding)
            except UnicodeDecodeError:
                continue
        return file_content.decode('utf-8', errors='replace')

    def _is_libro_compras_headers(self, header_values):
        return LIBRO_COMPRAS_SIGNATURE_HEADER in header_values

    def _read_csv_rows(self, file_content):
        text = self._decode_csv_text(file_content)
        rows = []
        for row in csv.reader(io.StringIO(text), delimiter=';'):
            rows.append(tuple(cell.strip() if cell not in (None, '') else None for cell in row))
        return rows

    @api.onchange('file_data')
    def _onchange_file_data(self):
        if not self.file_data:
            self.pre_analysis_message = False
            self.is_auto_detected = False
            return

        filename = (self.filename or '').lower()
        try:
            file_content = base64.b64decode(self.file_data)

            if filename.endswith('.csv'):
                rows = self._read_csv_rows(file_content)
                header_row = next((row for row in rows[:5] if row), None)
                if header_row and self._is_libro_compras_headers(header_row):
                    self.import_type = 'in_invoice'
                    self.pre_analysis_message = "Pre-análisis: Detectado 'Libro de Compras' (Compras)"
                    self.is_auto_detected = True
                else:
                    self.pre_analysis_message = "Pre-análisis: Tipo de comprobante no detectado automáticamente"
                    self.is_auto_detected = False
                return self._onchange_import_type()

            # Try to detect type from header
            # "Mis Comprobantes Recibidos" vs "Mis Comprobantes Emitidos"
            wb = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
            ws = wb.active
            # Get first cell value (A1)
            first_cell = ws.cell(row=1, column=1).value
            if first_cell and isinstance(first_cell, str):
                detected = False
                if 'Recibidos' in first_cell:
                    self.import_type = 'in_invoice'
                    self.pre_analysis_message = "Pre-análisis: Detectado 'Comprobantes Recibidos' (Proveedores)"
                    detected = True
                elif 'Emitidos' in first_cell:
                    self.import_type = 'out_invoice'
                    self.pre_analysis_message = "Pre-análisis: Detectado 'Comprobantes Emitidos' (Clientes)"
                    detected = True
                else:
                    self.pre_analysis_message = "Pre-análisis: Tipo de comprobante no detectado automáticamente"

                self.is_auto_detected = detected

                # Trigger journal update and domain set
                return self._onchange_import_type()
        except Exception:
            # Ignore errors here, let validation handle it
            self.pre_analysis_message = "Error en Pre-análisis: No se pudo leer el archivo"
            self.is_auto_detected = False

    @api.onchange('import_type')
    def _onchange_import_type(self):
        # Clear fields
        self.journal_purchase_id = False
        self.journal_sale_id = False

        company = self.env.company

        if self.import_type == 'out_invoice':
            domain = [('type', '=', 'sale'), ('company_id', '=',
                                              company.id), ('l10n_latam_use_documents', '=', True)]
            journal = self.env['account.journal'].search(domain, limit=1)
            if journal:
                self.journal_sale_id = journal
        else:
            domain = [('type', '=', 'purchase'), ('company_id', '=',
                                                  company.id), ('l10n_latam_use_documents', '=', True)]
            journal = self.env['account.journal'].search(domain, limit=1)
            if journal:
                self.journal_purchase_id = journal

    def _normalize_cuit(self, cuit):
        if not cuit:
            return False
        cuit = str(cuit)
        if cuit.endswith('.0'):
            cuit = cuit[:-2]
        return re.sub(r'\D', '', cuit)

    def _normalize_doc_code(self, raw_code):
        # El "Libro de Compras" trae el código AFIP puro (ej. "1", "11", "13"),
        # a veces con decimales residuales (ej. "1.0") si pasó por una planilla de cálculo.
        if raw_code is None:
            return False
        raw_code = str(raw_code).strip()
        if not raw_code:
            return False
        try:
            return str(int(float(raw_code)))
        except (TypeError, ValueError):
            return raw_code

    def _get_document_type(self, doc_name):
        # Mapeo básico de nombres ARCA a Document Types de Odoo
        # ARCA: "1 - Factura A", "6 - Factura B", "11 - Factura C"
        # Odoo l10n_ar: buscar por codigo
        code_match = re.match(r'^(\d+)\s-', doc_name)
        if code_match:
            code = code_match.group(1)
            # Buscar document type por codigo AFIP
            doc_type = self.env['l10n_latam.document.type'].search([
                ('code', '=', code),
                ('country_id.code', '=', 'AR')
            ], limit=1)
            return doc_type
        return False

    def _find_tax(self, amount, type_tax_use):
        # Búsqueda aproximada de impuestos por monto
        # ARCA columns: "IVA 21%", "IVA 10,5%", "IVA 27%"
        # amount es el valor porcentual (21.0, 10.5)
        # Buscar impuesto activo de ese tipo, de la compañía del wizard. Sin el
        # filtro de company_id, bajo un usuario con acceso a varias compañías
        # (ej. Superusuario, usado por el endpoint de importación automática,
        # que no aplica las reglas de registro que normalmente ocultan
        # impuestos de otras compañías) esto podía traer el impuesto de OTRA
        # compañía y romper la factura con "Incompatible companies".
        taxes = self.env['account.tax'].search([
            ('type_tax_use', '=', type_tax_use),
            ('amount', '=', amount),
            ('amount_type', '=', 'percent'),
            ('country_id.code', '=', 'AR'),
            ('company_id', '=', self.company_id.id),
            ('active', '=', True)
        ])
        return taxes[0] if taxes else False

    def _get_document_type_by_code(self, code):
        # El "Libro de Compras" ya trae el código AFIP puro (sin texto "N - Nombre"),
        # así que se busca directo por código, sin necesidad de parsear texto.
        if not code:
            return self.env['l10n_latam.document.type']
        return self.env['l10n_latam.document.type'].search([
            ('code', '=', code),
            ('country_id.code', '=', 'AR')
        ], limit=1)

    def _get_tax_by_xmlid_or_name(self, name, type_tax_use, xmlid_suffix=False):
        # Compartido entre el flujo legado (Mis Comprobantes) y el de Libro de Compras
        # para resolver impuestos de Exento/No Gravado/No Corresponde.
        tax = False
        if xmlid_suffix:
            xml_id = f"account.{self.company_id.id}_{xmlid_suffix}"
            tax = self.env.ref(xml_id, raise_if_not_found=False)
            if tax and tax.type_tax_use == type_tax_use:
                return tax

        return self.env['account.tax'].search([
            ('name', '=', name),
            ('type_tax_use', '=', type_tax_use),
            ('company_id', '=', self.company_id.id)
        ], limit=1)

    def _get_percepcion_tax(self, name_candidates, type_tax_use, ilike_fallback=None):
        # Busca un impuesto de percepción (Ganancias/IIBB) por nombre exacto,
        # probando cada candidato en 'name_candidates'. Si no encuentra nada
        # (ej. diferencias de espacios/formato en el nombre real cargado en
        # Odoo), cae a una búsqueda parcial ('ilike') con 'ilike_fallback'.
        base_domain = [
            ('company_id', '=', self.company_id.id),
            ('type_tax_use', '=', type_tax_use),
            ('active', '=', True),
        ]
        name_domain = ['|'] * (len(name_candidates) - 1) + [('name', '=', n) for n in name_candidates]
        tax = self.env['account.tax'].search(base_domain + name_domain, limit=1)
        if not tax and ilike_fallback:
            tax = self.env['account.tax'].search(base_domain + [('name', 'ilike', ilike_fallback)], limit=1)
        return tax

    def _get_percepcion_iva_tax(self, rate, type_tax_use):
        # 'Percepción de IVA' tiene dos variantes en el catálogo (ej. 3% para
        # netos gravados al 21%, 1,5% para netos al 10,5%). Se identifica por
        # la alícuota real del impuesto (amount) en vez del string exacto del
        # nombre, que es frágil ante diferencias de formato (ej. "Perc IVA
        # (3 %)" vs "Perc IVA (3%)" con espacio distinto antes del %).
        return self.env['account.tax'].search([
            ('company_id', '=', self.company_id.id),
            ('type_tax_use', '=', type_tax_use),
            ('active', '=', True),
            ('amount_type', '=', 'percent'),
            ('amount', '=', rate),
            ('name', 'ilike', 'IVA'),
        ], limit=1)

    def _force_percepcion_amount(self, move, tax, declared_amount):
        """ Sobrescribe el importe de la línea de impuesto que Odoo calculó
        automáticamente para una percepción (Ganancias/IIBB/Percepción IVA)
        adjunta a la línea de Neto Gravado IVA de mayor monto, para que
        coincida exactamente con el importe declarado por ARCA en el 'Libro
        de Compras' (que ARCA calcula sobre una base distinta a la de esa
        única línea, por lo que el % nominal del impuesto no reproduce ese
        valor). 'declared_amount' viene en la moneda del comprobante (ARS o
        USD si es en moneda extranjera). Ajusta la contrapartida (línea de
        Proveedores) para mantener el asiento balanceado. """
        tax_line = move.line_ids.filtered(lambda l: l.tax_line_id == tax)[:1]
        counterpart = move.line_ids.filtered(
            lambda l: l.account_id.account_type in ('liability_payable', 'asset_receivable'))[:1]
        if not tax_line or not counterpart:
            return

        def _debit_credit(balance):
            return {'debit': balance, 'credit': 0.0} if balance >= 0 else {'debit': 0.0, 'credit': -balance}

        rate = move.invoice_currency_rate or 1.0
        sign = 1 if tax_line.amount_currency >= 0 else -1
        new_amount_currency = sign * declared_amount
        new_balance = new_amount_currency / rate

        delta_balance = new_balance - tax_line.balance
        delta_amount_currency = new_amount_currency - tax_line.amount_currency
        new_counterpart_balance = counterpart.balance - delta_balance
        new_counterpart_amount_currency = counterpart.amount_currency - delta_amount_currency

        tax_line.with_context(check_move_validity=False).write(
            {**_debit_credit(new_balance), 'amount_currency': new_amount_currency})
        counterpart.with_context(check_move_validity=False).write(
            {**_debit_credit(new_counterpart_balance), 'amount_currency': new_counterpart_amount_currency})

        move._compute_amount()
        if hasattr(move, '_compute_tax_totals'):
            move._compute_tax_totals()

    def _find_existing_move(self, partner_name, line_move_type, doc_type, target_partners, pos_int, num_int, date):
        candidates = [
            f"{pos_int:05d}-{num_int:08d}",
            f"{pos_int:04d}-{num_int:08d}",
        ]

        def _check_duplicate_in_moves(moves):
            for move in moves:
                if move.l10n_latam_document_number in candidates:
                    return move
            return False

        move = False
        if target_partners:
            domain = [
                ('company_id', '=', self.company_id.id),
                ('move_type', '=', line_move_type),
                ('partner_id', 'in', target_partners.ids),
                ('l10n_latam_document_type_id', '=', doc_type.id),
            ]
            move = _check_duplicate_in_moves(self.env['account.move'].search(domain))

        if not move:
            domain_broad = [
                ('company_id', '=', self.company_id.id),
                ('move_type', '=', line_move_type),
                ('l10n_latam_document_type_id', '=', doc_type.id),
            ]
            if date:
                domain_broad.append(('invoice_date', '=', date))
            move = _check_duplicate_in_moves(self.env['account.move'].search(domain_broad))
            if move:
                _logger.info(f"DUPLICATE FOUND (BROAD): {partner_name} - MATCH: {move.name} on different partner {move.partner_id.name}")

        return move

    def _finalize_analysis(self, lines_values, unique_partners, company_currency):
        self.line_ids = [(5, 0, 0)] + [(0, 0, val) for val in lines_values]

        partner_filter_values = []
        for p_data in unique_partners.values():
            partner_filter_values.append((0, 0, {
                'name': p_data['name'],
                'cuit': p_data['cuit'],
                'partner_id': p_data['partner_id'],
                'invoice_count': p_data['invoice_count'],
                'amount_total': p_data['amount_total'],
                'currency_id': company_currency.id,
                'selected': True,  # Pre-select ALL by default
            }))
        self.partner_filter_ids = [(5, 0, 0)] + partner_filter_values

        self.state = 'review'
        return {
            'type': 'ir.actions.act_window',
            'name': 'Importación de Facturas de ARCA',
            'res_model': self._name,
            'view_mode': 'form',
            'res_id': self.id,
            'target': 'new',
        }

    def _analyze_libro_compras(self, rows, header_row_idx):
        """ Parsea el CSV 'Libro de Compras' de ARCA (ARCA > Libro de IVA Compras).
        A diferencia del Excel de 'Mis Comprobantes', este archivo:
          - Es exclusivamente de Compras (no trae datos de Ventas).
          - No incluye CAE.
          - Trae el código de comprobante AFIP puro (sin texto), un único
            'Número de Comprobante' (sin rango Desde/Hasta), y discrimina
            percepciones/otros tributos en columnas separadas en vez de un
            único monto "Otros Tributos".
        """
        self.import_type = 'in_invoice'

        header_row = rows[header_row_idx]
        header_values = [str(h).strip() if h is not None else None for h in header_row]
        missing_headers = [h for h in LIBRO_COMPRAS_REQUIRED_HEADERS if h not in header_values]
        if missing_headers:
            raise UserError(_(
                "El archivo 'Libro de Compras' no tiene la estructura esperada.\nColumnas faltantes: {}"
            ).format(', '.join(missing_headers)))

        headers = {}
        for col_idx, cell_val in enumerate(header_row):
            if cell_val:
                headers[str(cell_val).strip()] = col_idx

        if not self.journal_purchase_id:
            domain = [
                ('type', '=', 'purchase'), ('company_id', '=', self.company_id.id),
                ('l10n_latam_use_documents', '=', True)
            ]
            journal = self.env['account.journal'].search(domain, limit=1)
            if journal:
                self.journal_purchase_id = journal

        company_currency = self.env.company.currency_id
        type_tax_use = 'purchase'
        unique_partners = {}
        lines_values = []

        for row in rows[header_row_idx + 1:]:
            if not row:
                continue

            def val(name, _row=row):
                idx = headers.get(name)
                if idx is None or idx >= len(_row):
                    return None
                return _row[idx]

            def fval(name):
                # Las Notas de Crédito vienen con los importes en negativo en el
                # archivo; el signo del comprobante ya lo determina el 'Tipo de
                # Comprobante' (columna 2 -> is_refund_line/move_type), así que
                # todo importe numérico se toma siempre en valor absoluto.
                return abs(_safe_float(val(name)))

            if not val('Fecha de Emisión'):
                continue

            date_str = str(val('Fecha de Emisión')).strip()
            try:
                date = datetime.strptime(date_str, '%Y-%m-%d').date()
            except ValueError:
                try:
                    date = datetime.strptime(date_str, '%d/%m/%Y').date()
                except ValueError:
                    date = fields.Date.today()

            code = self._normalize_doc_code(val('Tipo de Comprobante'))
            doc_type = self._get_document_type_by_code(code)

            if doc_type:
                is_refund_line = doc_type.internal_type == 'credit_note'
                is_c_type = doc_type.l10n_ar_letter == 'C'
                prefix = doc_type.doc_code_prefix or f"COD{code}"
                display_name = (doc_type.name or '').strip()
                document_type_name = f"{code} - {display_name.title()}" if display_name else f"Comprobante Tipo {code}"
            else:
                is_refund_line = False
                is_c_type = False
                prefix = f"COD{code}" if code else "COD"
                document_type_name = f"Comprobante Tipo {code} (no reconocido)" if code else "Comprobante (tipo no reconocido)"

            line_move_type = 'in_refund' if is_refund_line else 'in_invoice'

            pos = val('Punto de Venta')
            number = val('Número de Comprobante')
            pos_int = _safe_int(pos)
            num_int = _safe_int(number)

            cuit = self._normalize_cuit(val('Nro. Doc. Vendedor'))
            partner_name = val('Denominación Vendedor') or 'Desconocido'

            amount_total = fval('Importe Total')
            moneda = (val('Moneda Original') or '').strip().upper()

            status = 'ready'
            error_msgs = []

            journal = self.journal_purchase_id
            if not journal:
                status = 'error'
                error_msgs.append("Falta seleccionar Diario de Compras.")

            # --- Moneda extranjera (USD/DOL) ---
            # El 'Libro de Compras' expresa los importes en la moneda original del
            # comprobante (no en pesos), e informa el tipo de cambio de ARCA (pesos
            # por dólar) en la columna 'Tipo de Cambio'. El resto de las columnas
            # numéricas de la fila quedan entonces en esa misma moneda.
            invoice_currency = company_currency
            invoice_currency_rate = 1.0
            if moneda in ('DOL', 'USD', 'U$S', 'U$'):
                usd_currency = self.env['res.currency'].with_context(active_test=False).search(
                    [('name', '=', 'USD')], limit=1)
                tipo_cambio = fval('Tipo de Cambio')
                if not usd_currency:
                    status = 'error'
                    error_msgs.append("No se encontró la moneda 'USD' en el sistema.")
                elif not usd_currency.active:
                    status = 'error'
                    error_msgs.append("La moneda 'USD' está inactiva. Actívela en Contabilidad > Configuración > Monedas.")
                elif not tipo_cambio:
                    status = 'error'
                    error_msgs.append("Comprobante en USD sin 'Tipo de Cambio' informado en el archivo.")
                else:
                    invoice_currency = usd_currency
                    # invoice_currency_rate en Odoo convierte de moneda de compañía (ARS)
                    # a moneda del comprobante (USD): 1 ARS = (1/tipo_cambio) USD.
                    invoice_currency_rate = 1.0 / tipo_cambio
            elif moneda not in ('PES', '$', 'ARS', ''):
                status = 'error'
                error_msgs.append(
                    f"Comprobante en moneda extranjera ({moneda or '?'}). No soportado en esta versión, requiere carga manual.")

            partners_with_cuit = self.env['res.partner'].search([
                ('vat', '=', cuit),
                ('parent_id', '=', False)
            ])
            partner = partners_with_cuit[0] if partners_with_cuit else False

            if not partner and not cuit:
                status = 'error'
                error_msgs.append("Falta CUIT para identificar contacto.")

            if not doc_type:
                status = 'error'
                error_msgs.append(f"Tipo de comprobante '{code}' no reconocido en el sistema.")

            move = False
            if doc_type and journal:
                target_partners = partners_with_cuit if partners_with_cuit else (partner if partner else False)
                move = self._find_existing_move(
                    partner_name, line_move_type, doc_type, target_partners, pos_int, num_int, date)
                if move:
                    status = 'exists'
                    msg_extra = ""
                    if partner and move.partner_id != partner:
                        msg_extra = f" (En contacto: {move.partner_id.name})"
                    error_msgs.append(
                        f"<b>Factura ya existe{msg_extra}</b>: <a href='#' data-oe-model='account.move' data-oe-id='{move.id}'>{move.name}</a> (Doc: {move.l10n_latam_document_number})")

            # --- Líneas de factura: base gravada / no gravada / exenta ---
            invoice_lines_data = []
            declared_components_total = 0.0
            # rate -> (índice en invoice_lines_data, monto neto), solo para tasas
            # con importe > 0. Se usa para asociar las percepciones (Ganancias,
            # IIBB, Percepción IVA) a la línea de Neto Gravado de mayor monto.
            rate_line_index = {}

            for rate, neto_col, iva_col in LIBRO_COMPRAS_TAX_RATE_COLUMNS:
                amount_neto = fval(neto_col)
                declared_components_total += amount_neto
                declared_iva_amount = fval(iva_col) if iva_col else 0.0
                if iva_col:
                    declared_components_total += declared_iva_amount
                if not amount_neto:
                    continue
                tax = self._find_tax(rate, type_tax_use)
                if not tax:
                    status = 'error'
                    error_msgs.append(f"No se encontró Impuesto {rate:g}% {type_tax_use}.")
                    continue

                # --- Corroboración: IVA calculado por Odoo vs. declarado por ARCA ---
                # Tolerancia de $3: en datos reales aparecen diferencias de
                # centavos por redondeo propio de ARCA (no siempre coincide
                # con el cálculo estándar de Odoo sobre el neto), sin ser un
                # error real del comprobante.
                if iva_col:
                    computed = tax.compute_all(amount_neto, currency=invoice_currency, quantity=1.0)
                    computed_iva_amount = sum(t['amount'] for t in computed['taxes'])
                    if abs(computed_iva_amount - declared_iva_amount) > 3.0:
                        status = 'error'
                        error_msgs.append(
                            "El IVA {:g}% calculado (${:,.2f}) sobre el Neto Gravado (${:,.2f}) no coincide con "
                            "el declarado en el archivo (${:,.2f}, diferencia ${:,.2f}). Verifique el comprobante en ARCA."
                            .format(rate, computed_iva_amount, amount_neto, declared_iva_amount,
                                    computed_iva_amount - declared_iva_amount))

                invoice_lines_data.append({
                    'price_unit': amount_neto,
                    'tax_ids': [tax.id],
                    'name': f'Importe Gravado {rate:g}%'
                })
                rate_line_index[rate] = (len(invoice_lines_data) - 1, amount_neto)

            amount_ng = fval('Importe No Gravado')
            if amount_ng:
                tax_ng = self._get_tax_by_xmlid_or_name(
                    'IVA No Gravado', type_tax_use, 'ri_tax_vat_no_gravado_compras')
                if not tax_ng:
                    tax_ng = self.env['account.tax'].search([
                        ('name', 'ilike', 'No Grav'),
                        ('type_tax_use', '=', type_tax_use),
                        ('company_id', '=', self.company_id.id)
                    ], limit=1)
                invoice_lines_data.append({
                    'price_unit': amount_ng,
                    'tax_ids': [tax_ng.id] if tax_ng else [],
                    'name': 'Conceptos No Gravados'
                })

            amount_ex = fval('Importe Exento')
            if amount_ex:
                tax_ex = self._get_tax_by_xmlid_or_name(
                    'IVA Exento', type_tax_use, 'ri_tax_vat_exento_compras')
                if not tax_ex:
                    tax_ex = self.env['account.tax'].search([
                        ('name', 'ilike', 'Exento'),
                        ('type_tax_use', '=', type_tax_use),
                        ('company_id', '=', self.company_id.id)
                    ], limit=1)
                invoice_lines_data.append({
                    'price_unit': amount_ex,
                    'tax_ids': [tax_ex.id] if tax_ex else [],
                    'name': 'Operaciones Exentas'
                })

            # Antes de agregar Municipales/Internos/percepciones: si a esta altura
            # no hay ninguna línea de base gravada/no gravada/exenta, es un
            # comprobante sin discriminar (típ. Factura C monotributista) y se
            # resuelve más abajo con el remanente como línea única.
            has_base_lines = bool(invoice_lines_data)

            # --- Percepciones / Otros Tributos discriminados ---
            # Se guardan en la línea para revisión y se agregan como líneas de
            # factura recién en action_import (usando el impuesto/cuenta que se
            # configure en el wizard). 'Crédito Fiscal Computable' queda afuera:
            # es informativo, no es un importe adicional del comprobante.
            tribute_amounts = {}
            tribute_total = 0.0
            for field_name, col_name, _label in LIBRO_COMPRAS_OTHER_TRIBUTES_COLUMNS:
                amount = fval(col_name)
                tribute_amounts[field_name] = amount
                tribute_total += amount

            credito_fiscal_computable = fval('Crédito Fiscal Computable')

            # --- Impuestos Municipales / Internos: mismo tratamiento que un Neto
            # Gravado IVA (línea propia, sin impuesto asociado, misma cuenta
            # contable que las líneas de Neto Gravado: no se fija account_id acá,
            # así hereda el override del wizard o la cuenta por defecto). ---
            if tribute_amounts['imp_municipales']:
                invoice_lines_data.append({
                    'price_unit': tribute_amounts['imp_municipales'],
                    'tax_ids': [],
                    'name': 'Impuestos Municipales'
                })
            if tribute_amounts['imp_internos']:
                invoice_lines_data.append({
                    'price_unit': tribute_amounts['imp_internos'],
                    'tax_ids': [],
                    'name': 'Impuestos Internos'
                })

            # --- Percepciones (Ganancias / IIBB / IVA): se asocian como impuesto
            # adicional en la línea de Neto Gravado IVA de mayor monto (una sola
            # línea, aunque haya varias alícuotas). El importe calculado por Odoo
            # sobre esa única línea normalmente NO coincide con el declarado por
            # ARCA (que lo calcula sobre otra base), así que se fuerza el importe
            # exacto sobre la línea de impuesto ya creada, recién en action_import
            # (una vez creado el asiento). Acá solo se resuelve el impuesto a
            # aplicar y se agrega su id al tax_ids de la línea elegida.
            forced_tax_amounts = []
            general_max_rate = max(rate_line_index, key=lambda r: rate_line_index[r][1]) if rate_line_index else None

            def _attach_percepcion(target_rate, tax, amount, label):
                if amount <= 0:
                    return
                if not tax:
                    nonlocal status
                    status = 'error'
                    error_msgs.append(f"No se encontró el impuesto de percepción '{label}'.")
                    return
                if target_rate is not None and target_rate in rate_line_index:
                    idx, _ = rate_line_index[target_rate]
                    invoice_lines_data[idx]['tax_ids'].append(tax.id)
                    forced_tax_amounts.append([tax.id, amount])
                else:
                    # Sin línea de Neto Gravado a la que asociarla (ej. Factura C
                    # sin discriminar): se agrega como línea propia, best-effort.
                    invoice_lines_data.append({
                        'price_unit': amount,
                        'tax_ids': [tax.id] if tax else [],
                        'name': label,
                    })

            if tribute_amounts['perc_otros_imp_nacionales']:
                tax_gan = self._get_percepcion_tax(TAX_NAME_PERC_GANANCIAS, type_tax_use, ilike_fallback='Gananc')
                _attach_percepcion(general_max_rate, tax_gan, tribute_amounts['perc_otros_imp_nacionales'], 'Percepción de Ganancias')

            if tribute_amounts['perc_iibb']:
                tax_iibb = self._get_percepcion_tax(TAX_NAME_PERC_IIBB, type_tax_use, ilike_fallback='IIBB')
                _attach_percepcion(general_max_rate, tax_iibb, tribute_amounts['perc_iibb'], 'Percepción de IIBB Mendoza')

            perc_iva_total = tribute_amounts['perc_iva'] + tribute_amounts['otros_tributos']
            if perc_iva_total:
                if general_max_rate == 10.5:
                    perc_iva_rate, perc_iva_target = 1.5, 10.5
                else:
                    perc_iva_rate = 3.0
                    perc_iva_target = 21.0 if 21.0 in rate_line_index else general_max_rate
                tax_perc_iva = self._get_percepcion_iva_tax(perc_iva_rate, type_tax_use)
                _attach_percepcion(perc_iva_target, tax_perc_iva, perc_iva_total, 'Percepción de IVA')

            # --- Reconciliación del Importe Total contra sus componentes discriminados ---
            # ARCA a veces exporta comprobantes donde el Total no coincide con la suma de
            # Neto/IVA por tasa + No Gravado + Exento + percepciones/tributos (se detectó
            # en datos reales). Ante esa inconsistencia se marca error en vez de crear una
            # factura con un monto que no se puede sustentar en el detalle del archivo.
            declared_components_total += amount_ng + amount_ex + tribute_total
            # Si el archivo no discrimina nada (típ. Factura C monotributista), no hay
            # nada que reconciliar: se resuelve más abajo con el remanente como línea única.
            if amount_total and declared_components_total > 0.01 and abs(amount_total - declared_components_total) > 0.10:
                status = 'error'
                error_msgs.append(
                    "El Importe Total (${:,.2f}) no coincide con la suma de sus componentes discriminados "
                    "del archivo (${:,.2f}, diferencia ${:,.2f}). Verifique el comprobante en ARCA."
                    .format(amount_total, declared_components_total, amount_total - declared_components_total))

            # Comprobantes sin desglose de gravado/no gravado/exento (típ. Factura C):
            # se registra el remanente (Total - percepciones/tributos) como línea única.
            if not has_base_lines and amount_total:
                remainder = amount_total - tribute_total
                if is_c_type:
                    tax_nc = self._get_tax_by_xmlid_or_name(
                        'IVA No Corresponde', type_tax_use, 'ri_tax_vat_no_corresponde_compras')
                    invoice_lines_data.append({
                        'price_unit': remainder,
                        'tax_ids': [tax_nc.id] if tax_nc else [],
                        'name': 'Importe General (Factura C)'
                    })
                else:
                    invoice_lines_data.append({
                        'price_unit': remainder,
                        'tax_ids': [],
                        'name': 'Importe General'
                    })

            invoice_vals = {
                'ref': f"{document_type_name} {pos}-{number}",
                'move_type': line_move_type,
                'invoice_date': date.strftime('%Y-%m-%d') if date else False,
                'partner_id': partner.id if partner else False,
                'journal_id': journal.id if journal else False,
                'l10n_latam_document_type_id': doc_type.id if doc_type else False,
                'l10n_latam_document_number': f"{pos_int:05d}-{num_int:08d}",
                'currency_id': invoice_currency.id,
                'invoice_currency_rate': invoice_currency_rate,
                'invoice_line_ids': [(0, 0, line) for line in invoice_lines_data],
                '_forced_tax_amounts': forced_tax_amounts,
            }

            try:
                preview_str = f"{prefix} {pos_int:05d}-{num_int:08d}"
            except Exception:
                preview_str = f"{document_type_name} {pos}-{number}"

            lines_values.append({
                'date': date,
                'document_type_name': document_type_name,
                'point_of_sale': pos,
                'number': number,
                'cuit': cuit,
                'afip_auth_code': False,
                'partner_name': partner_name,
                'partner_id': partner.id if partner else False,
                'journal_id': journal.id if journal else False,
                'amount_total': amount_total,
                'currency_id': invoice_currency.id,
                'status': status,
                'error_desc': '<br/>'.join(error_msgs) if error_msgs else False,
                'invoice_values': json.dumps(invoice_vals),
                'move_id': move.id if move else False,
                'to_import': True if status == 'ready' else False,
                'preview_desc': preview_str,
                'is_refund': is_refund_line,
                'source_format': 'libro_compras',
                'credito_fiscal_computable': credito_fiscal_computable,
                'otros_tributos': tribute_amounts['otros_tributos'],
                'perc_otros_imp_nacionales': tribute_amounts['perc_otros_imp_nacionales'],
                'perc_iibb': tribute_amounts['perc_iibb'],
                'imp_municipales': tribute_amounts['imp_municipales'],
                'perc_iva': tribute_amounts['perc_iva'],
                'imp_internos': tribute_amounts['imp_internos'],
                'iva_21_amount': fval('Importe IVA 21%'),
                'iva_105_amount': fval('Importe IVA 10,5%'),
            })

            if partner_name not in unique_partners:
                unique_partners[partner_name] = {
                    'name': partner_name,
                    'cuit': cuit,
                    'partner_id': partner.id if partner else False,
                    'invoice_count': 0,
                    'amount_total': 0.0,
                }
            unique_partners[partner_name]['invoice_count'] += 1
            unique_partners[partner_name]['amount_total'] += amount_total

        return self._finalize_analysis(lines_values, unique_partners, company_currency)

    def action_create_partners(self):
        """ Batch create partners for selected (or all) missing ones """
        self.ensure_one()
        
        # Determine which filters to process:
        # Strict: process only SELECTED missing ones
        selected = self.partner_filter_ids.filtered('selected')
        to_process = selected
        
        if not to_process:
             # Should not happen given UI button visibility, but safety first
             return {'type': 'ir.actions.act_window_close'}
        
        # Filter down to those that actually need creation
        to_process = to_process.filtered(lambda p: not p.partner_id)
        
        cuit_type = self.env['l10n_latam.identification.type'].search([('name', '=', 'CUIT')], limit=1)
        
        for p_filter in to_process:
            # Double check existence (maybe created by another user)
            domain = [('vat', '=', p_filter.cuit)] if p_filter.cuit else [('name', '=', p_filter.name)]
            partner = self.env['res.partner'].search(domain, limit=1)
            
            if not partner:
                vals = {
                    'name': p_filter.name,
                    'company_type': 'company',
                    'l10n_ar_afip_responsibility_type_id': self.env.ref('l10n_ar.res_IVARI').id,
                }
                if p_filter.cuit:
                    vals['vat'] = p_filter.cuit
                    if cuit_type:
                        vals['l10n_latam_identification_type_id'] = cuit_type.id
                        
                partner = self.env['res.partner'].create(vals)
                partner.message_post(body=Markup(_("Partner dado de alta automáticamente por <b>Importación de Facturas ARCA</b>.")))
            
            # Link back to filter
            p_filter.partner_id = partner.id
            
        # Refresh lines to link new partners
        # We can iterate lines and update partner_id where missing
        cuit_partner_map = {p.cuit: p.partner_id.id for p in self.partner_filter_ids if p.cuit and p.partner_id}
        name_partner_map = {p.name: p.partner_id.id for p in self.partner_filter_ids if p.name and p.partner_id}
        
        for line in self.line_ids:
            if not line.partner_id:
                if line.cuit and line.cuit in cuit_partner_map:
                    line.partner_id = cuit_partner_map[line.cuit]
                elif line.partner_name in name_partner_map:
                    line.partner_id = name_partner_map[line.partner_name]

        return {
            'type': 'ir.actions.act_window',
            'name': 'Importación de Comprobantes',
            'res_model': self._name,
            'view_mode': 'form',
            'res_id': self.id,
            'target': 'new',
        }




    def action_analyze(self):
        self.ensure_one()
        if not self.file_data:
            raise UserError(_("Por favor suba un archivo Excel."))

        # Initialize dictionary to track unique partners for the filter tab
        unique_partners = {}

        try:
            # Decode file
            file_content = base64.b64decode(self.file_data)
        except Exception as e:
            raise UserError(
                _("No se pudo decodificar el archivo. Asegúrese de que sea un archivo válido. Error: %s") % e)

        filename = (self.filename or '').lower()

        # 0. Validate File Extension
        if not filename.endswith('.xlsx') and not filename.endswith('.csv'):
            raise UserError(
                _("Formato de archivo inválido. Solo se permiten archivos Excel (.xlsx) o CSV (.csv) de ARCA."))

        # --- Formato CSV: por ahora solo se reconoce el diseño "Libro de Compras" ---
        if filename.endswith('.csv'):
            rows = self._read_csv_rows(file_content)
            header_row_idx = next(
                (i for i, row in enumerate(rows) if row and self._is_libro_compras_headers(row)), None)
            if header_row_idx is None:
                raise UserError(_(
                    "El archivo CSV no coincide con el formato 'Libro de Compras' esperado.\n"
                    "Verifique que sea el archivo descargado desde ARCA > Libro de IVA Compras (Compras)."
                ))
            return self._analyze_libro_compras(rows, header_row_idx)

        # --- Formato Excel "Mis Comprobantes" (legado) ---
        if 'openpyxl' not in globals():
            raise UserError(
                _("La librería 'openpyxl' no está instalada. Por favor contacte al administrador."))

        wb = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
        ws = wb.active  # Asumimos primera hoja

        # Headers mapping (based on provided sample)
        # We need to map column index to field
        headers = {}
        header_row_idx = 1

        rows = list(ws.iter_rows(values_only=True))

        start_row = 0

        # Expected Headers for Strict Validation based on Type
        # Common headers
        ARCA_EXPECTED_HEADERS_COMMON = [
            'Fecha', 'Tipo', 'Punto de Venta', 'Número Desde', 'Número Hasta', 'Cód. Autorización',
            'Tipo Cambio', 'Moneda', 'Imp. Total'
        ]

        # Specific headers
        if self.import_type == 'in_invoice':
            # Recibidos -> Emisor info
            ARCA_EXPECTED_HEADERS_SPECIFIC = [
                'Tipo Doc. Emisor', 'Nro. Doc. Emisor', 'Denominación Emisor']
        else:
            # Emitidos -> Receptor info
            ARCA_EXPECTED_HEADERS_SPECIFIC = [
                'Tipo Doc. Receptor', 'Nro. Doc. Receptor', 'Denominación Receptor']

        ARCA_EXPECTED_HEADERS = ARCA_EXPECTED_HEADERS_COMMON + ARCA_EXPECTED_HEADERS_SPECIFIC

        found_header_row = False
        for i, row in enumerate(rows):
            if row and 'Fecha' in row and 'Tipo' in row:
                start_row = i + 1
                found_header_row = True

                # --- STRICT STRUCTURE VALIDATION ---
                row_values = [str(cell).strip()
                              for cell in row if cell is not None]

                # Check for missing headers
                missing_headers = [
                    h for h in ARCA_EXPECTED_HEADERS if h not in row_values]

                if missing_headers:
                    raise UserError(_("El archivo no tiene la estructura esperada de ARCA para '{}'.\nColumnas faltantes: {}")
                                    .format(dict(self._fields['import_type'].selection).get(self.import_type), ', '.join(missing_headers)))

                # Map headers
                for col_idx, cell_val in enumerate(row):
                    if cell_val:
                        headers[cell_val] = col_idx
                break

        if not found_header_row:
            raise UserError(
                _("No se encontró la fila de encabezados válida (Fecha, Tipo, etc) o el formato no coincide con el de ARCA."))

        # --- VALIDACIÓN CUIT COMPAÑÍA ---
        # Fila 0, Celda 0 usualmente contiene: "Mis Comprobantes ... - CUIT 30718858514"
        try:
            # Re-leer fila 0 para asegurar
            first_row = next(ws.iter_rows(
                min_row=1, max_row=1, values_only=True))
            header_title = first_row[0]
            if header_title and 'CUIT' in str(header_title):
                # Extraer CUIT (últimos dígitos o regex)
                cuit_match = re.search(r'CUIT\s+(\d+)', str(header_title))
                if cuit_match:
                    excel_cuit = cuit_match.group(1)
                    # Normalizar Company CUIT: Remover AR, guiones y espacios
                    company_vat = self.env.company.vat or ''
                    # A veces l10n_ar agrega prefix AR or similiar
                    company_cuit = re.sub(r'\D', '', company_vat)

                    if company_cuit and excel_cuit != company_cuit:
                        raise UserError(_(
                            "El CUIT del archivo ({}) no coincide con el CUIT de la compañía actual ({}).\n"
                            "Por favor verifique que está importando el archivo correcto."
                        ).format(excel_cuit, company_cuit))
        except UserError:
            raise
        except Exception as e:
            _logger.warning(f"No se pudo validar el CUIT del encabezado: {e}")
            # No bloqueamos si falla el parseo del título, pero logueamos.

        lines_values = []

        # Pre-fetch generic data
        company_currency = self.env.company.currency_id

        for row_idx in range(start_row, len(rows)):
            row = rows[row_idx]
            if not row[headers.get('Fecha')]:
                continue

            # Extraer Datos Básicos
            date_str = row[headers['Fecha']]
            # Parse Date dd/mm/yyyy
            try:
                date = datetime.strptime(
                    date_str, '%d/%m/%Y').date() if isinstance(date_str, str) else date_str
            except:
                date = fields.Date.today()

            doc_type_name = row[headers.get('Tipo', -1)]
            pos = row[headers.get('Punto de Venta', -1)]
            number_from = row[headers.get('Número Desde', -1)]
            # number_to = row[headers.get('Número Hasta')]

            # --- Detección de Nota de Crédito ---
            is_refund_line = doc_type_name and 'Nota de Crédito' in str(doc_type_name)
            line_move_type = self.import_type
            if is_refund_line:
                line_move_type = 'in_refund' if self.import_type == 'in_invoice' else 'out_refund'

            cuit_col = 'Nro. Doc. Emisor' if self.import_type == 'in_invoice' else 'Nro. Doc. Receptor'
            name_col = 'Denominación Emisor' if self.import_type == 'in_invoice' else 'Denominación Receptor'

            # Use safe get just in case logic fails or col missing despite validation
            cuit_raw = row[headers.get(cuit_col)] if headers.get(
                cuit_col) is not None else False
            cuit = self._normalize_cuit(cuit_raw)
            partner_name = row[headers.get(name_col)] if headers.get(
                name_col) is not None else 'Desconocido'

            # CAE Extraction
            afip_auth_code = False
            for head_alias in ['Cód. Autorización', 'CAE', 'C.A.E', 'Codigo Autorizacion', 'Código Autorización']:
                if headers.get(head_alias) is not None:
                    afip_auth_code = row[headers[head_alias]]
                    break

            amount_total = row[headers.get('Imp. Total', -1)]

            # --- Otros Tributos Detection ---
            def _parse_float_safe(val):
                try:
                    if isinstance(val, (int, float)):
                        return float(val)
                    if isinstance(val, str):
                        return float(val.replace(',', '.'))
                    return 0.0
                except:
                    return 0.0

            otros_tributos_raw = row[headers.get('Otros Tributos', -1)] if headers.get('Otros Tributos') is not None else 0.0
            otros_tributos_amount = _parse_float_safe(otros_tributos_raw)
            
            iva_21_raw = row[headers.get('IVA 21%', -1)] if headers.get('IVA 21%') is not None else 0.0
            iva_21_extracted_amount = _parse_float_safe(iva_21_raw)
            
            iva_105_raw = row[headers.get('IVA 10,5%', -1)] if headers.get('IVA 10,5%') is not None else 0.0
            iva_105_extracted_amount = _parse_float_safe(iva_105_raw)

            # --- VALIDACIONES Y RESOLUCIONES ---
            status = 'ready'
            error_msgs = []

            # 1. Diario
            journal = False
            if self.import_type == 'in_invoice':
                journal = self.journal_purchase_id
                if not journal:
                    status = 'error'
                    error_msgs.append("Falta seleccionar Diario de Compras.")
            else:
                # Ventas:
                # User requires usage of the VALIDATED journal in the wizard, not auto-search logic.
                if self.journal_sale_id:
                    journal = self.journal_sale_id
                else:
                    status = 'error'
                    error_msgs.append(f"No se seleccionó Diario de Venta.")

            # 2. Partner (Contacto)
            # Find ALL partners with this CUIT to check for duplicates later
            partners_with_cuit = self.env['res.partner'].search([
                ('vat', '=', cuit),
                ('parent_id', '=', False)
            ])

            # For assignment, we prefer the first one found or create new if none
            partner = partners_with_cuit[0] if partners_with_cuit else False

            if not partner and not cuit:
                status = 'error'
                error_msgs.append("Falta CUIT para identificar contacto.")

            # 3. Documento Existente
            doc_type = self._get_document_type(doc_type_name)
            move = False

            # Ensure headers exist
            if headers.get('Número Desde') is None:
                # Try 'Número' if 'Número Desde' missing (some variations)
                if headers.get('Número') is not None:
                    headers['Número Desde'] = headers['Número']
                else:
                    raise UserError(
                        "No se encuentra columna 'Número Desde' o 'Número'")

            if doc_type and journal:
                # 0. Validate Point of Sale matching for Sales (FACTURAS DE VENTA)
                if self.import_type == 'out_invoice' and journal:
                    # Check attribute existence
                    journal_pos = getattr(
                        journal, 'l10n_ar_afip_pos_number', None)

                    _logger.info(
                        f"VALIDATION DEBUG: File POS={pos}, Journal={journal.name}, Journal POS={journal_pos}")

                    if journal_pos is not None:
                        try:
                            file_pos = int(pos) if pos else 0
                            if file_pos != journal_pos:
                                status = 'error'
                                error_msgs.append(
                                    f"Punto de Venta incorrecto. Archivo: {file_pos}, Diario: {journal_pos}")
                        except Exception as e:
                            _logger.error(f"Error parsing/validating POS: {e}")
                    else:
                        _logger.warning(
                            "Journal has no l10n_ar_afip_pos_number set or field missing")

                # Robust Duplicate Check
                def _parse_int_safe(val):
                    try:
                        if isinstance(val, (int, float)):
                            return int(val)
                        if isinstance(val, str):
                            # Handle "123.0"
                            if val.endswith('.0'):
                                val = val[:-2]
                            # Remove thousands separators and keep only digits
                            val = val.replace('.', '').replace(',', '')
                            return int(re.sub(r'\D', '', val) or 0)
                        return 0
                    except:
                        return 0

                pos_int = _parse_int_safe(pos)
                num_int = _parse_int_safe(number_from)
                
                # --- DUPLICATE DETECTION DEBUG ---
                # Log critical values to understand why detection might fail
                _logger.info(f"--- ANALYZING DUPLICATE: {partner_name} ---")
                _logger.info(f"Raw POS: {pos}, Raw Num: {number_from}")
                _logger.info(f"Parsed POS: {pos_int}, Parsed Num: {num_int}")

                # Formatos posibles de número de documento en Odoo (AR)
                # Standard: 00002-00001234 (5 pos, 8 number)
                # Legacy: 0002-00001234 (4 pos)
                candidates = [
                    f"{pos_int:05d}-{num_int:08d}",
                    f"{pos_int:04d}-{num_int:08d}",
                ]
                _logger.info(f"Candidates: {candidates}")

                # Check against ALL partners with same CUIT if available, otherwise just general check?
                # Actually, duplicate check should be against the specific partner(s) to avoid false positives with other partners having same number?
                target_partners = partners_with_cuit if partners_with_cuit else (
                    partner if partner else False)

                def _check_duplicate_in_moves(moves, candidates):
                    for move in moves:
                        if move.l10n_latam_document_number in candidates:
                            return move
                    return False

                if target_partners:
                    # Search by partner + doc type (stored fields)
                    domain = [
                        ('company_id', '=', self.company_id.id),
                        ('move_type', '=', line_move_type),
                        ('partner_id', 'in', target_partners.ids),
                        ('l10n_latam_document_type_id', '=', doc_type.id),
                    ]
                    # Fetch potential matches
                    possible_moves = self.env['account.move'].search(domain)
                    move = _check_duplicate_in_moves(possible_moves, candidates)

                # FALLBACK: Check broadly (across other partners) if not found specific
                if not move:
                    domain_broad = [
                        ('company_id', '=', self.company_id.id),
                        ('move_type', '=', line_move_type),
                        ('l10n_latam_document_type_id', '=', doc_type.id),
                        # Rely on Python filtering for the number
                    ]
                    # Warning: This could be large, but we need to find if it exists on ANY partner
                    # Optimization: Use ref? 'ref' might contain the number but format varies.
                    # Let's trust that for a specific DocType + Company, list isn't infinite, 
                    # OR we can try to search by 'ref' if feasible, but 'ref' in my wizard 
                    # includes "Factura A ...". Standard Odoo might differ.
                    # Safest: Search generic but maybe limit to recent? Or just rely on limit?
                    # No, limit won't help if the duplicate is 50th.
                    # Better fallback: Don't do the broad search if it risks performance, 
                    # OR assume if it wasn't found on the partner, it might not be critical to block?
                    # User needs to know if it exists on ANOTHER partner.
                    # Optimisation: Filter by date if available?
                    
                    if date:
                         domain_broad.append(('invoice_date', '=', date))
                    
                    possible_moves_broad = self.env['account.move'].search(domain_broad)
                    move = _check_duplicate_in_moves(possible_moves_broad, candidates)
                    
                    if move:
                         _logger.info(f"DUPLICATE FOUND (BROAD): {partner_name} - MATCH: {move.name} on different partner {move.partner_id.name}")

                if move:
                    status = 'exists'
                    # Enhanced Error Message for User Debugging
                    msg_extra = ""
                    if partner and move.partner_id != partner:
                        msg_extra = f" (En contacto: {move.partner_id.name})"

                    error_msgs.append(
                        f"<b>Factura ya existe{msg_extra}</b>: <a href='#' data-oe-model='account.move' data-oe-id='{move.id}'>{move.name}</a> (Doc: {move.l10n_latam_document_number})")

            # 4. Impuestos (Recopilar líneas)
            tax_lines = []

            # --- Initialize Unique Partners Dict ---
            if 'unique_partners' not in locals():
                unique_partners = {}
            # Mapeo de columnas de impuetos ARCA a tasas
            tax_mapping = {
                'IVA 21%': 21.0,
                'IVA 10,5%': 10.5,
                'IVA 27%': 27.0,
                'IVA 5%': 5.0,
                'IVA 2,5%': 2.5,
                # 'IVA 0%': 0.0 # Exento or 0?
            }

            # Type Tax Use verification
            type_tax_use = 'purchase' if self.import_type == 'in_invoice' else 'sale'

            # Stores (tax_id, base_amount) or product line
            invoice_lines_data = []

            # Netos
            # ARCA has specific logic. Typically: "Neto Grav. IVA 21%", "IVA 21%"
            # We can create lines based on Netos

            has_taxes = False
            for col_name, rate in tax_mapping.items():
                neto_col = f"Neto Grav. {col_name}" if col_name != 'IVA 0%' else 'Neto Grav. IVA 0%'
                # Sometimes headers change slightly "Neto Grav. IVA 21%" matches sample

                if headers.get(neto_col) is not None and row[headers[neto_col]]:
                    amount_neto = row[headers[neto_col]]
                    if amount_neto:
                        tax = self._find_tax(rate, type_tax_use)
                        if not tax:
                            status = 'error'
                            error_msgs.append(
                                f"No se encontró Impuesto {rate}% {type_tax_use}.")
                        else:
                            has_taxes = True
                            invoice_lines_data.append({
                                'price_unit': amount_neto,
                                'tax_ids': [tax.id],
                                'name': f'Importe Gravado {rate}%'
                            })

            # Helper to find tax by name or XmlID
            def _get_tax_by_xmlid_or_name(name, type_tax_use, xmlid_suffix=False):
                tax = False

                # 1. Try XmlID if suffix provided (e.g. 'ri_tax_vat_no_gravado_ventas')
                if xmlid_suffix:
                    # Construct XmlID: account.{company_id}_{suffix}
                    # Example: account.5_ri_tax_vat_no_gravado_ventas
                    # Note: We need to search by external ID which is 'module.name'
                    # The module is 'account'. The name part includes company_id.

                    xml_id = f"account.{self.company_id.id}_{xmlid_suffix}"
                    tax = self.env.ref(xml_id, raise_if_not_found=False)
                    if tax and tax.type_tax_use == type_tax_use:
                        return tax

                # 2. Fallback to Name Search
                return self.env['account.tax'].search([
                    ('name', '=', name),
                    ('type_tax_use', '=', type_tax_use),
                    ('company_id', '=', self.company_id.id)
                ], limit=1)

            # Exentos / No Gravados
            if headers.get('Neto No Gravado') is not None and row[headers['Neto No Gravado']]:
                amount_ng = row[headers['Neto No Gravado']]
                if amount_ng:
                    # Try to find specific tax for No Gravado using XmlID first
                    suffix = 'ri_tax_vat_no_gravado_compras' if type_tax_use == 'purchase' else 'ri_tax_vat_no_gravado_ventas'
                    tax_ng = _get_tax_by_xmlid_or_name(
                        'IVA No Gravado', type_tax_use, suffix)

                    if not tax_ng:
                        # Fallback to search by description or similar
                        tax_ng = self.env['account.tax'].search([
                            ('name', 'ilike', 'No Grav'),
                            ('type_tax_use', '=', type_tax_use),
                            ('company_id', '=', self.company_id.id)
                        ], limit=1)

                    invoice_lines_data.append({
                        'price_unit': amount_ng,
                        'tax_ids': [tax_ng.id] if tax_ng else [],
                        'name': 'Conceptos No Gravados'
                    })

            # Factura C / B logic (Monotributo / Consumidor Final often have no split taxes)
            # If no taxes found yet, and we have a total amount, check if we need to apply "No Corresponde" or "Exento"
            # Simple check for Factura C, Nota de Credito C, etc.
            is_c_type = ' C' in doc_type_name

            if not has_taxes and not invoice_lines_data:
                # Check for explicit "Exento" column first
                if headers.get('Op. Exentas') is not None and row[headers['Op. Exentas']]:
                    amount_ex = row[headers['Op. Exentas']]
                    if amount_ex:
                        suffix = 'ri_tax_vat_exento_compras' if type_tax_use == 'purchase' else 'ri_tax_vat_exento_ventas'
                        tax_ex = _get_tax_by_xmlid_or_name(
                            'IVA Exento', type_tax_use, suffix)
                        if not tax_ex:
                            tax_ex = self.env['account.tax'].search([
                                ('name', 'ilike', 'Exento'),
                                ('type_tax_use', '=', type_tax_use),
                                ('company_id', '=', self.company_id.id)
                            ], limit=1)

                        invoice_lines_data.append({
                            'price_unit': amount_ex,
                            'tax_ids': [tax_ex.id] if tax_ex else [],
                            'name': 'Operaciones Exentas'
                        })

                # If still no lines and we have total, maybe it's Factura C
                if not invoice_lines_data and amount_total:
                    if is_c_type:
                        # Apply IVA No Corresponde
                        suffix = 'ri_tax_vat_no_corresponde_compras' if type_tax_use == 'purchase' else 'ri_tax_vat_no_corresponde_ventas'
                        tax_nc = _get_tax_by_xmlid_or_name(
                            'IVA No Corresponde', type_tax_use, suffix)
                        invoice_lines_data.append({
                            'price_unit': amount_total,
                            'tax_ids': [tax_nc.id] if tax_nc else [],
                        })

            # Construct JSON for creation
            invoice_vals = {
                'ref': f"{doc_type_name} {pos}-{number_from}",
                'move_type': line_move_type,
                'invoice_date': date.strftime('%Y-%m-%d') if date else False,
                # Will need creation logic if False
                'partner_id': partner.id if partner else False,
                'journal_id': journal.id if journal else False,
                'l10n_latam_document_type_id': doc_type.id if doc_type else False,
                'l10n_latam_document_number': f"{int(pos):05d}-{int(number_from):08d}",
                'invoice_line_ids': [(0, 0, line) for line in invoice_lines_data]
            }

            if not invoice_lines_data and status == 'ready':
                # Fallback if no specific tax columns found but total exists (Simpler inv?)
                if amount_total:
                    invoice_vals['invoice_line_ids'] = [(0, 0, {
                        'price_unit': amount_total,
                        'tax_ids': [],
                        'name': 'Importe General'
                    })]

            # Generate Preview String (Invoice Number)
            try:
                # Clean doc type name (e.g. "001 - Factura A" -> "Factura A")
                clean_doc_name = doc_type_name
                if ' - ' in doc_type_name:
                    clean_doc_name = doc_type_name.split(' - ')[1]

                # Abbreviation Mapping
                abbr_map = {
                    'Factura A': 'FA-A', 'Factura B': 'FA-B', 'Factura C': 'FA-C',
                    'Nota de Crédito A': 'NC-A', 'Nota de Crédito B': 'NC-B', 'Nota de Crédito C': 'NC-C',
                    'Nota de Débito A': 'ND-A', 'Nota de Débito B': 'ND-B', 'Nota de Débito C': 'ND-C',
                    'Recibo A': 'REC-A', 'Recibo B': 'REC-B', 'Recibo C': 'REC-C',
                    'Recibo X': 'REC-X',
                }

                prefix = abbr_map.get(clean_doc_name, clean_doc_name)

                preview_str = f"{prefix} {int(pos):05d}-{int(number_from):08d}"
            except Exception as e:
                preview_str = f"{doc_type_name} {pos}-{number_from}"

            lines_values.append({
                'date': date,
                'document_type_name': doc_type_name,
                'point_of_sale': pos,
                'number': number_from,
                'cuit': cuit,
                'afip_auth_code': afip_auth_code,
                'partner_name': partner_name,
                'partner_id': partner.id if partner else False,
                'journal_id': journal.id if journal else False,
                'amount_total': amount_total,
                'currency_id': company_currency.id,
                'status': status,
                'error_desc': '<br/>'.join(error_msgs) if error_msgs else False,
                'invoice_values': json.dumps(invoice_vals),
                'move_id': move.id if move else False,
                'to_import': True if status == 'ready' else False,
                'preview_desc': preview_str,
                'is_refund': is_refund_line,
                'otros_tributos': otros_tributos_amount,
                'iva_21_amount': iva_21_extracted_amount,
                'iva_105_amount': iva_105_extracted_amount
            })

            # --- Partner Filter Aggregation ---
            if partner_name not in unique_partners:
                unique_partners[partner_name] = {
                    'name': partner_name,
                    'cuit': cuit,
                    'partner_id': partner.id if partner else False,
                    'invoice_count': 0,
                    'amount_total': 0.0,
                }
            unique_partners[partner_name]['invoice_count'] += 1
            unique_partners[partner_name]['amount_total'] += amount_total

        self.line_ids = [(5, 0, 0)] + [(0, 0, val) for val in lines_values]

        # Populate Partner Filter
        partner_filter_values = []
        for p_name, p_data in unique_partners.items():
            is_missing = not p_data['partner_id']
            partner_filter_values.append((0, 0, {
                'name': p_data['name'],
                'cuit': p_data['cuit'],
                'partner_id': p_data['partner_id'],
                'invoice_count': p_data['invoice_count'],
                'amount_total': p_data['amount_total'],
                'currency_id': company_currency.id,
                'selected': True,  # Pre-select ALL by default
            }))
        self.partner_filter_ids = [(5, 0, 0)] + partner_filter_values

        self.state = 'review'
        return {
            'type': 'ir.actions.act_window',
            'name': 'Importación de Facturas de ARCA',
            'res_model': self._name,
            'view_mode': 'form',
            'res_id': self.id,
            'target': 'new',
        }

    def action_back(self):
        self.ensure_one()
        self.state = 'upload'
        return {
            'type': 'ir.actions.act_window',
            'name': 'Importación de Facturas de ARCA',
            'res_model': self._name,
            'view_mode': 'form',
            'res_id': self.id,
            'target': 'new',
        }

    def action_import(self):
        self.ensure_one()
        created_moves = self.env['account.move']

        # Determine lines to process based on Partner Filter
        lines_to_process = self.line_ids
        selected_partners = self.partner_filter_ids.filtered('selected')
        if selected_partners:
            selected_names = selected_partners.mapped('name')
            lines_to_process = lines_to_process.filtered(
                lambda l: l.partner_name in selected_names)

        for line in lines_to_process:
            if not line.to_import or line.status != 'ready':
                continue

            vals = json.loads(line.invoice_values)
            forced_tax_amounts = vals.pop('_forced_tax_amounts', [])

            # 1. Partner Creation if needed
            exist_partners = self.env['res.partner'].search(
                [('vat', '=', line.cuit)], limit=1)
            if not vals.get('partner_id') and not exist_partners:
                # Search for CUIT identification type
                cuit_type = self.env['l10n_latam.identification.type'].search(
                    [('name', '=', 'CUIT')], limit=1)

                # Create partner
                partner = self.env['res.partner'].create({
                    'name': line.partner_name,
                    'vat': line.cuit,
                    'company_type': 'company',  # Force Company
                    'l10n_latam_identification_type_id': cuit_type.id if cuit_type else False,
                    # Default to Responsable Inscripto logic might need review but ok for now
                    'l10n_ar_afip_responsibility_type_id': self.env.ref('l10n_ar.res_IVARI').id,
                })
                # Chatter message for partner
                partner.message_post(body=Markup(
                    _("Partner dado de alta automáticamente por <b>Importación de Facturas ARCA</b>.")))

                vals['partner_id'] = partner.id
            elif exist_partners:
                vals['partner_id'] = exist_partners[0].id
            # 2. Create Move
            try:
                # If we have AFIP Auth Code, we must ensure auth_mode is CAE (especially for vendor bills or offline mode)
                if line.afip_auth_code:
                    vals['l10n_ar_afip_auth_mode'] = 'CAE'

                # Apply Analytic Distribution if selected
                if self.analytic_distribution and vals.get('invoice_line_ids'):
                    for cmd in vals['invoice_line_ids']:
                        if len(cmd) == 3 and isinstance(cmd[2], dict):
                            cmd[2]['analytic_distribution'] = self.analytic_distribution

                # Apply Account ID if selected (Override)
                if self.account_id and vals.get('invoice_line_ids'):
                    for cmd in vals['invoice_line_ids']:
                        if len(cmd) == 3 and isinstance(cmd[2], dict):
                            cmd[2]['account_id'] = self.account_id.id

                # --- OTROS TRIBUTOS LOGIC (solo formato legado 'Mis Comprobantes') ---
                # Se genera una línea de factura por cada percepción/tributo discriminado
                # presente. Todas usan el impuesto/cuenta configurado en 'Otros Tributos'
                # del wizard. Para 'Libro de Compras' este esquema ya NO se usa: Municipales
                # e Internos quedan resueltos como líneas propias desde el análisis, y
                # Ganancias/IIBB/Percepción IVA se fuerzan sobre la línea de Neto Gravado
                # de mayor monto más abajo (ver 'forced_tax_amounts').
                tribute_values = []
                otros_tributos_tax_tags = self.env['account.account.tag']
                if line.source_format != 'libro_compras':
                    tribute_values = [
                        (label, getattr(line, fname))
                        for fname, _col, label in LIBRO_COMPRAS_OTHER_TRIBUTES_COLUMNS
                        if getattr(line, fname) > 0
                    ]
                    if tribute_values and self.tax_other_tributes_id:
                        rep_lines = self.tax_other_tributes_id.invoice_repartition_line_ids if vals.get('move_type') in ('in_invoice', 'out_invoice') else self.tax_other_tributes_id.refund_repartition_line_ids
                        tax_rep = rep_lines.filtered(lambda r: r.repartition_type == 'tax')
                        tax_account_id = tax_rep[0].account_id.id if tax_rep and tax_rep[0].account_id else (self.account_id.id or False)
                        otros_tributos_tax_tags = tax_rep[0].tag_ids if tax_rep else self.env['account.account.tag']

                        # Fetch exempt tax (IVA No Corresponde) to avoid AFIP validation errors
                        type_tax_use = 'purchase' if vals.get('move_type') in ('in_invoice', 'in_refund') else 'sale'

                        suffix = 'ri_tax_vat_no_corresponde_compras' if type_tax_use == 'purchase' else 'ri_tax_vat_no_corresponde_ventas'
                        xml_id = f"account.{self.env.company.id}_{suffix}"
                        tax_nc = self.env.ref(xml_id, raise_if_not_found=False)

                        if not tax_nc:
                            tax_nc = self.env['account.tax'].search([
                                ('company_id', '=', self.env.company.id),
                                ('type_tax_use', '=', type_tax_use),
                                '|', ('name', '=', 'IVA No Corresponde'), ('name', 'ilike', 'No Corresp')
                            ], limit=1)

                        if not vals.get('invoice_line_ids'):
                            vals['invoice_line_ids'] = []

                        for label, amount in tribute_values:
                            vals['invoice_line_ids'].append((0, 0, {
                                'name': label,
                                'quantity': 1,
                                'price_unit': amount,
                                'account_id': tax_account_id,
                                'tax_ids': [(6, 0, tax_nc.ids)] if tax_nc else [],
                                'analytic_distribution': self.analytic_distribution or False,
                            }))

                move = self.env['account.move'].create(vals)

                # Append Perception Tax Grids post-creation (formato legado)
                if tribute_values and self.tax_other_tributes_id and otros_tributos_tax_tags:
                    tribute_labels = {label for label, _amount in tribute_values}
                    ot_lines = move.invoice_line_ids.filtered(lambda l: l.name in tribute_labels)
                    if ot_lines:
                        # Append tags to respect base tags set by tax_nc
                        ot_lines.with_context(check_move_validity=False).write({
                            'tax_tag_ids': [(4, t.id) for t in otros_tributos_tax_tags]
                        })

                # --- Percepciones 'Libro de Compras' (Ganancias/IIBB/Percepción IVA) ---
                # Se fuerza el importe exacto declarado por ARCA en la línea de impuesto
                # ya generada por Odoo sobre la línea de Neto Gravado de mayor monto.
                if line.source_format == 'libro_compras' and forced_tax_amounts:
                    for tax_id, forced_amount in forced_tax_amounts:
                        forced_tax = self.env['account.tax'].browse(int(tax_id))
                        self._force_percepcion_amount(move, forced_tax, float(forced_amount))

                # --- CHATTER & CAE LOGIC ---
                msg_body = Markup(
                    _("Factura importada desde archivo ARCA: <b>%s</b>")) % (self.filename or 'Desconocido')
                move.message_post(body=msg_body)

                # Check AFIP Verification Type logic correctly using the new computed mode
                move_model = self.env['account.move']
                has_enterprise = 'l10n_ar_afip_auth_mode' in move_model._fields
                has_community = 'afip_auth_mode' in move_model._fields

                if self.cae_status_mode == 'available' and line.afip_auth_code:
                    import datetime
                    cae_due_date = line.date + datetime.timedelta(days=10) if line.date else False
                    
                    if has_enterprise:
                        move.with_context(check_move_validity=False).write({
                            'l10n_ar_afip_auth_code': line.afip_auth_code,
                            'l10n_ar_afip_auth_code_due': cae_due_date,
                            'l10n_ar_afip_result': 'A',  # Always A for ARCA history
                        })
                    elif has_community:
                        move.with_context(check_move_validity=False).write({
                            'afip_auth_code': line.afip_auth_code,
                            'afip_auth_code_due': cae_due_date,
                            'afip_result': 'A',  # Always A for ARCA history
                        })

                created_moves |= move
                line.status = 'exists'
                line.move_id = move.id
                line.to_import = False
            except Exception as e:
                line.status = 'error'
                line.error_desc = str(e)

        # --- CREATE HISTORY RECORD ---
        if created_moves:
            self.env['l10n_ar.arca.import.history'].create({
                'name': f"Importación {fields.Datetime.now().strftime('%d/%m/%Y %H:%M')} - {self.filename or 'Sin Nombre'}",
                'filename': self.filename,
                'file_data': self.file_data,
                'import_type': self.import_type,
                'move_ids': [(6, 0, created_moves.ids)]
            })
            
        # CLEAR FIELDS FOR NEXT BATCH
        self.analytic_distribution = False
        self.account_id = False
        self.tax_other_tributes_id = False

        if len(created_moves) > 0:
            title = "Importación Completada"
            message = f"Se importaron {len(created_moves)} comprobantes exitosamente."
            type_notif = 'success'
        else:
            title = "Importación Finalizada"
            message = "Ningún comprobante fue importado. Verifique si hay errores (EJ: falta Cuenta Contable)."
            type_notif = 'warning'

        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': title,
                'message': message,
                'type': type_notif,
                'sticky': False,
                'next': {
                    'type': 'ir.actions.act_window',
                    'name': 'Importación de Facturas de ARCA',
                    'res_model': self._name,
                    'view_mode': 'form',
                    'views': [(False, 'form')],
                    'res_id': self.id,
                    'target': 'new',
                }
            }
        }


class L10nArImportArcaLine(models.TransientModel):
    _name = 'l10n_ar.arca.import.line'
    _description = 'Línea de Importación ARCA'
    _order = 'date desc, number desc'

    wizard_id = fields.Many2one('l10n_ar.arca.import.wizard', string='Wizard')

    date = fields.Date(string='Fecha')
    document_type_name = fields.Char(string='Tipo Doc.')
    point_of_sale = fields.Char(string='Pto. Venta')
    number = fields.Char(string='Número')
    cuit = fields.Char(string='CUIT')

    partner_name = fields.Char(string='Razón Social')
    partner_id = fields.Many2one('res.partner', string='Partner')

    amount_total = fields.Monetary(
        string='Total', currency_field='currency_id')
    currency_id = fields.Many2one('res.currency', string='Moneda')

    status = fields.Selection([
        ('ready', 'Listo'),
        ('exists', 'Ya Existe'),
        ('error', 'Error')
    ], string='Estado', default='ready')

    error_desc = fields.Html(string='Detalle Error')

    # Store JSON data for invoice creation
    invoice_values = fields.Text(string='Valores JSON')

    move_id = fields.Many2one('account.move', string='Factura Creada')
    journal_id = fields.Many2one('account.journal', string='Diario')

    afip_auth_code = fields.Char(string='CAE')

    to_import = fields.Boolean(string='Importar', default=False)

    preview_desc = fields.Char(string='Vista Previa')
    is_refund = fields.Boolean(string='Es NC')
    otros_tributos = fields.Float(string='Otros Tributos')
    iva_21_amount = fields.Float(string='IVA 21%')
    iva_105_amount = fields.Float(string='IVA 10.5%')

    source_format = fields.Selection([
        ('mis_comprobantes', 'Mis Comprobantes (Excel)'),
        ('libro_compras', 'Libro de Compras (CSV)'),
    ], string='Origen del Dato', default='mis_comprobantes')

    # Percepciones e impuestos discriminados (solo presentes en 'Libro de Compras')
    credito_fiscal_computable = fields.Float(
        string='Crédito Fiscal Computable',
        help="Informativo: porción del IVA total que ARCA considera computable como crédito fiscal. No se suma al total del comprobante.")
    perc_otros_imp_nacionales = fields.Float(string='Perc./Pagos Cta. Otros Imp. Nac.')
    perc_iibb = fields.Float(string='Percepciones IIBB')
    imp_municipales = fields.Float(string='Impuestos Municipales')
    perc_iva = fields.Float(string='Perc./Pagos Cta. IVA')
    imp_internos = fields.Float(string='Impuestos Internos')

    def action_toggle_import(self):
        for line in self:
            if line.status == 'ready':
                line.to_import = not line.to_import

        # Reload the wizard view to reflect changes
        # Use wizard_id to target the correct record
        return {
            'type': 'ir.actions.act_window',
            'name': 'Importación de Facturas de ARCA',
            'res_model': 'l10n_ar.arca.import.wizard',
            'res_id': self[0].wizard_id.id,
            'view_mode': 'form',
            'target': 'new',
        }
